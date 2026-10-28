"""The exercise library and the user's workspace.

The library ships as plain CSV files in data/ (easy to edit and review in pull requests). On first run
it is built into an Excel workbook (FitView.xlsx) in the workspace, which the user can then edit freely.
"""

import csv
import shutil
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .planfile import DAYS, PLAN_HEADERS, UNITS, parse_amount, write_plan_file

REPO = Path(__file__).resolve().parent.parent
DATA_DIR = REPO / "data"

EXERCISE_BASE = ["Exercise", "Type", "Category", "Equipment", "Knee Load", "Default Sets", "Default Amount", "Unit", "Per Side", "Notes"]
TYPES = ["Gym", "Calisthenics"]
CATEGORIES = ["Push", "Pull", "Legs", "Core", "Skill", "Grip", "Cardio", "Mobility"]
KNEE_LOADS = ["Low", "Medium", "High"]
WEIGHTS = {1.0, 0.5, 0.25}

# key -> (default, allowed range, description written next to it in the Settings sheet)
SETTINGS = {
    "stabiliser_weight": (0.0, (0.0, 0.25), "Credit for stabiliser muscles (0.25 weight in Exercises). 0 = not counted (research style), 0.25 = counted."),
    "emom_factor": (0.5, (0.25, 1.0), "How much one EMOM set counts compared with a hard straight set (0.5 = 10 EMOM sets count as 5)."),
}

README_LINES = [
    "FitView library workbook",
    "",
    "Exercises  - the exercise library. One column per muscle. Weight: 1 = main mover, 0.5 = secondary, 0.25 = stabiliser. Leave blank for none.",
    "             Default Amount is numbers only (10 or 8-12); Unit is reps / s / min / m / steps / rounds; Per Side = yes when done per side.",
    "             Add rows freely. Do not rename the muscle column headers (they must match the Muscles sheet).",
    "Muscles    - the muscle list the body map uses. Reference only - do not edit IDs.",
    "Settings   - counting settings (also changeable in the app).",
    "Plans      - each weekly plan is its own file in the Plans folder (file name = plan title).",
    "             Columns: Day, Session, Exercise, Sets, Amount, Unit, Per Side, Load, Load Type (kg / +kg / -kg / band, blank = bodyweight), EMOM, Hidden, Notes.",
    "             Deleted plans are moved to Plans/_deleted, never erased. Exports (PDF / images) are saved next to the plan files.",
    "",
    "Close a file in Excel before saving changes to it from the FitView app (Excel locks open files).",
]


class LibraryError(ValueError):
    """The shipped CSV data is inconsistent (bad weight, unknown muscle, duplicate name...)."""


def _read_csv(path):
    with open(path, newline="", encoding="utf-8-sig") as f:
        return list(csv.reader(f))


def load_muscles(data_dir=DATA_DIR):
    rows = _read_csv(data_dir / "muscles.csv")
    return [tuple(r[:4]) for r in rows[1:] if r and r[0].strip()]


def load_exercises(data_dir=DATA_DIR):
    """Returns (muscle column names, rows). Each row: EXERCISE_BASE values followed by one weight per muscle column.
    Raises LibraryError if the data is inconsistent."""
    muscles = load_muscles(data_dir)
    names = {m[1] for m in muscles}
    rows = _read_csv(data_dir / "exercises.csv")
    header = [h.strip() for h in rows[0]]
    if header[: len(EXERCISE_BASE)] != EXERCISE_BASE:
        raise LibraryError(f"exercises.csv must start with columns {EXERCISE_BASE}")
    muscle_cols = header[len(EXERCISE_BASE) :]
    unknown = [m for m in muscle_cols if m not in names]
    if unknown:
        raise LibraryError(f"exercises.csv has muscle columns not in muscles.csv: {unknown}")

    out, seen, problems = [], set(), []
    for line, r in enumerate(rows[1:], start=2):
        if not r or not r[0].strip():
            continue
        r = [c.strip() for c in r] + [""] * (len(header) - len(r))
        name = r[0]
        if name.lower() in seen:
            problems.append(f"line {line}: duplicate exercise {name!r}")
        seen.add(name.lower())
        if r[1] not in TYPES:
            problems.append(f"line {line}: {name}: Type must be one of {TYPES}")
        if r[2] and r[2] not in CATEGORIES:
            problems.append(f"line {line}: {name}: Category must be one of {CATEGORIES}")
        if r[4] and r[4] not in KNEE_LOADS:
            problems.append(f"line {line}: {name}: Knee Load must be blank or one of {KNEE_LOADS}")
        if r[7] and r[7] not in UNITS:
            problems.append(f"line {line}: {name}: Unit must be one of {UNITS}")
        if r[6] and parse_amount(r[6]) == (None, None):
            problems.append(f"line {line}: {name}: Default Amount must be a number or range like 8-12")
        weights = []
        for col, v in zip(muscle_cols, r[len(EXERCISE_BASE) :]):
            if not v:
                weights.append(None)
                continue
            try:
                w = float(v)
            except ValueError:
                w = None
            if w not in WEIGHTS:
                problems.append(f"line {line}: {name}: weight for {col!r} must be 1, 0.5 or 0.25 (got {v!r})")
            weights.append(w)
        if not any(weights):
            problems.append(f"line {line}: {name}: needs at least one muscle weight")
        base = r[: len(EXERCISE_BASE)]
        base[5] = int(base[5]) if base[5].isdigit() else (base[5] or None)
        lo, hi = parse_amount(base[6])
        base[6] = lo if hi is None else (f"{lo}-{hi}" if lo is not None else None)
        out.append([v if v != "" else None for v in base] + weights)
    if problems:
        raise LibraryError("Problems in exercises.csv:\n  " + "\n  ".join(problems))
    return muscle_cols, out


def load_example_plans(data_dir=DATA_DIR):
    """{title: rows} from data/examples/*.csv. Rows are in PLAN_HEADERS order."""
    plans = {}
    for path in sorted((data_dir / "examples").glob("*.csv")):
        rows = _read_csv(path)
        header = [h.strip() for h in rows[0]]
        if header != PLAN_HEADERS:
            raise LibraryError(f"{path.name}: header must be {PLAN_HEADERS}")
        body = []
        for r in rows[1:]:
            if not any(c.strip() for c in r):
                continue
            r = [c.strip() for c in r] + [""] * (len(header) - len(r))
            if r[0] not in DAYS:
                raise LibraryError(f"{path.name}: unknown day {r[0]!r}")
            r[3] = int(r[3]) if r[3].isdigit() else (r[3] or None)
            body.append([c if c != "" else None for c in r])
        plans[path.stem] = body
    return plans


# ---------------------------------------------------------------- workbook
HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def _style_header(ws, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "B2"


def write_exercises_sheet(wb, rows, muscle_names, index=None):
    """rows: EXERCISE_BASE values followed by one weight per muscle column."""
    ws = wb.create_sheet("Exercises", index)
    base = EXERCISE_BASE
    ws.append(base + list(muscle_names))
    for r in rows:
        ws.append(r)
    _style_header(ws, len(base) + len(muscle_names))
    ws.row_dimensions[1].height = 60
    for c in range(len(base) + 1, len(base) + len(muscle_names) + 1):
        ws.cell(row=1, column=c).alignment = Alignment(text_rotation=90, horizontal="center")
        ws.column_dimensions[get_column_letter(c)].width = 5
    for col, width in zip("ABCDEFGHIJ", (44, 13, 10, 16, 10, 8, 10, 8, 8, 40)):
        ws.column_dimensions[col].width = width
    last = max(ws.max_row, 2)
    end = last + 200
    first_m, last_m = get_column_letter(len(base) + 1), get_column_letter(len(base) + len(muscle_names))
    rng = f"{first_m}2:{last_m}{end}"
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=["1"], fill=PatternFill("solid", fgColor="E5484D")))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=["0.5"], fill=PatternFill("solid", fgColor="F5A35C")))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=["0.25"], fill=PatternFill("solid", fgColor="FCE3B0")))
    lists = (
        ('"' + ",".join(TYPES) + '"', f"B2:B{end}"),
        ('"' + ",".join(CATEGORIES) + '"', f"C2:C{end}"),
        ('"' + ",".join(KNEE_LOADS) + '"', f"E2:E{end}"),
        ('"' + ",".join(UNITS) + '"', f"H2:H{end}"),
        ('"yes"', f"I2:I{end}"),
        ('"1,0.5,0.25"', rng),
    )
    for formula, ref in lists:
        dv = DataValidation(type="list", formula1=formula, allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(ref)
    ws.auto_filter.ref = f"A1:{last_m}{last}"
    return ws


def write_settings_sheet(wb, values=None):
    values = values or {k: v[0] for k, v in SETTINGS.items()}
    if "Settings" in wb.sheetnames:
        del wb["Settings"]
    ws = wb.create_sheet("Settings")
    ws.append(["Key", "Value", "Description"])
    for key, (_, _, desc) in SETTINGS.items():
        ws.append([key, values[key], desc])
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["C"].width = 110
    return ws


def build_library_workbook(path, data_dir=DATA_DIR):
    """Creates FitView.xlsx (README, Muscles, Exercises, Settings) from the CSV data."""
    muscles = load_muscles(data_dir)
    muscle_cols, rows = load_exercises(data_dir)
    wb = Workbook()
    ws = wb.active
    ws.title = "README"
    for line in README_LINES:
        ws.append([line])
    ws["A1"].font = Font(bold=True, size=14)
    ws.column_dimensions["A"].width = 150

    ws = wb.create_sheet("Muscles")
    ws.append(["ID", "Name", "Group", "Drawn on"])
    for m in muscles:
        ws.append(list(m))
    _style_header(ws, 4)
    for col, width in zip("ABCD", (12, 28, 12, 12)):
        ws.column_dimensions[col].width = width

    write_exercises_sheet(wb, rows, muscle_cols)
    write_settings_sheet(wb)
    wb.save(path)


def ensure_workspace(home, data_dir=DATA_DIR):
    """First run: create the workspace with the library workbook and the example plans.
    Never overwrites existing files. Returns a list of what was created."""
    home = Path(home)
    plans = home / "Plans"
    plans.mkdir(parents=True, exist_ok=True)
    created = []
    xlsx = home / "FitView.xlsx"
    if not xlsx.exists():
        build_library_workbook(xlsx, data_dir)
        created.append(xlsx.name)
    if not any(p for p in plans.glob("*.xlsx") if not p.name.startswith("~$")):
        library = {r[0].lower() for r in load_exercises(data_dir)[1]}
        for title, rows in load_example_plans(data_dir).items():
            unknown = sorted({r[2] for r in rows if r[2] and r[2].lower() not in library})
            if unknown:
                raise LibraryError(f"Example plan {title!r} uses exercises not in the library: {unknown}")
            write_plan_file(plans / f"{title}.xlsx", rows)
            created.append(f"Plans/{title}.xlsx")
    return created


def reset_library(home, data_dir=DATA_DIR):
    """Rebuilds FitView.xlsx from the shipped CSVs, keeping a backup of the old one."""
    xlsx = Path(home) / "FitView.xlsx"
    if xlsx.exists():
        shutil.copy2(xlsx, xlsx.with_name("FitView.backup.xlsx"))
    build_library_workbook(xlsx, data_dir)
