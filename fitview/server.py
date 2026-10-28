"""FitView local server: serves the UI, reads the exercise library (FitView.xlsx in the workspace)
and reads/writes one .xlsx file per plan in <workspace>/Plans.

Run:  python -m fitview            (see `python -m fitview --help`)
"""

import argparse
import base64
import json
import mimetypes
import os
import platform
import shutil
import subprocess
import sys
import tempfile
import threading
import time
import uuid
import webbrowser
from datetime import datetime
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, quote, urlparse

from openpyxl import load_workbook

from . import __version__
from .library import REPO, SETTINGS, LibraryError, ensure_workspace, write_settings_sheet
from .planfile import DAYS, LOAD_TYPES, PLAN_HEADERS, UNITS, format_amount, parse_amount, read_plan_rows, write_plan_file

STATIC = Path(__file__).resolve().parent / "static"
DEFAULT_HOME = Path(os.environ.get("FITVIEW_HOME") or REPO / "workspace")
HOST = "127.0.0.1"
DEFAULT_PORT = 8765

BASE_COLS = ["Exercise", "Type", "Category", "Equipment", "Knee Load", "Default Sets", "Default Amount", "Unit", "Per Side", "Notes"]
BAD_NAME_CHARS = set('<>:"/\\|?*')
WINDOWS_RESERVED = {"con", "prn", "aux", "nul", *(f"com{i}" for i in range(1, 10)), *(f"lpt{i}" for i in range(1, 10))}
LOCK = threading.Lock()
PRINT_DAYS = {}  # token -> plan days for an in-progress PDF export (lets unsaved edits be exported)

# Set by configure(); module-level so request handlers can use them.
HOME = XLSX = PLANS_DIR = DELETED_DIR = None
PORT = DEFAULT_PORT


def configure(home):
    """Points the server at a workspace folder (library workbook + Plans/)."""
    global HOME, XLSX, PLANS_DIR, DELETED_DIR
    HOME = Path(home)
    XLSX = HOME / "FitView.xlsx"
    PLANS_DIR = HOME / "Plans"
    DELETED_DIR = PLANS_DIR / "_deleted"


configure(DEFAULT_HOME)


class ApiError(Exception):
    def __init__(self, status, message):
        super().__init__(message)
        self.status = status


def _str(v):
    return "" if v is None else str(v).strip()


def _num(v):
    if v in (None, ""):
        return None
    try:
        f = float(v)
    except (TypeError, ValueError):
        return None
    return int(f) if f.is_integer() else f


def _unit(v):
    v = _str(v).lower()
    return v if v in UNITS else "reps"


def _load_type(v):
    v = _str(v).lower().replace("−", "-")
    return v if v in LOAD_TYPES else ""


def _yes(v):
    return _str(v).lower() in ("yes", "y", "true", "1", "x")


def locked(fn):
    """Runs a file operation, turning Excel's file lock into a readable error."""
    try:
        return fn()
    except PermissionError as e:
        name = Path(e.filename).name if e.filename else "the file"
        raise ApiError(423, f"{name} is locked - close it (Excel / PDF viewer) and try again") from None


# ---------------------------------------------------------------- library
def read_library():
    wb = load_workbook(XLSX, read_only=True)
    try:
        muscles = []
        for row in wb["Muscles"].iter_rows(min_row=2, values_only=True):
            if row and row[0]:
                muscles.append({"id": _str(row[0]), "name": _str(row[1]), "group": _str(row[2]), "drawn": _str(row[3])})
        by_name = {m["name"].lower(): m["id"] for m in muscles}

        rows = wb["Exercises"].iter_rows(values_only=True)
        header = [_str(h) for h in next(rows)]
        col = {h: i for i, h in enumerate(header)}
        muscle_cols = [(i, by_name[h.lower()]) for i, h in enumerate(header) if h.lower() in by_name]
        unknown_headers = [h for h in header if h and h not in BASE_COLS and h.lower() not in by_name]

        exercises = []
        for row in rows:
            get = lambda h, row=row: _str(row[col[h]]) if h in col and col[h] < len(row) else ""
            name = get("Exercise")
            if not name:
                continue
            weights = {}
            for i, mid in muscle_cols:
                w = _num(row[i]) if i < len(row) else None
                if w:
                    weights[mid] = w
            lo, hi = parse_amount(row[col["Default Amount"]] if "Default Amount" in col else None)
            exercises.append(
                {
                    "name": name,
                    "type": get("Type"),
                    "category": get("Category"),
                    "equipment": get("Equipment"),
                    "knee": get("Knee Load"),
                    "defaultSets": _num(get("Default Sets")),
                    "defaultAmountMin": lo,
                    "defaultAmountMax": hi,
                    "unit": _unit(get("Unit")),
                    "perSide": _yes(get("Per Side")),
                    "notes": get("Notes"),
                    "muscles": weights,
                }
            )
        return muscles, exercises, unknown_headers
    finally:
        wb.close()


# ---------------------------------------------------------------- settings (Settings sheet in FitView.xlsx)
def read_settings():
    out = {k: v[0] for k, v in SETTINGS.items()}
    wb = load_workbook(XLSX, read_only=True)
    try:
        if "Settings" in wb.sheetnames:
            for row in wb["Settings"].iter_rows(min_row=2, values_only=True):
                key = _str(row[0] if row else "")
                if key in SETTINGS and _num(row[1]) is not None:
                    lo, hi = SETTINGS[key][1]
                    out[key] = min(hi, max(lo, float(row[1])))
    finally:
        wb.close()
    return out


def write_settings(values):
    current = read_settings()
    for key, (_, (lo, hi), _) in SETTINGS.items():
        v = _num(values.get(key))
        if v is not None:
            if not lo <= v <= hi:
                raise ApiError(400, f"{key} must be between {lo} and {hi}")
            current[key] = v
    wb = load_workbook(XLSX)
    write_settings_sheet(wb, current)
    locked(lambda: wb.save(XLSX))
    return current


# ---------------------------------------------------------------- plans (one file each)
def plan_names():
    PLANS_DIR.mkdir(parents=True, exist_ok=True)
    return sorted((p.stem for p in PLANS_DIR.glob("*.xlsx") if not p.name.startswith("~$")), key=str.lower)


def plan_path(name):
    return PLANS_DIR / f"{name}.xlsx"


def find_plan(name):
    """Case-insensitive lookup (Windows and macOS file names are case-insensitive)."""
    for n in plan_names():
        if n.lower() == _str(name).lower():
            return n
    return None


def validate_title(title):
    title = _str(title)
    if not title:
        raise ApiError(400, "Plan title is required")
    if len(title) > 80:
        raise ApiError(400, "Plan title must be 80 characters or fewer")
    if any(c in BAD_NAME_CHARS or ord(c) < 32 for c in title) or title.endswith("."):
        raise ApiError(400, 'Plan title cannot contain < > : " / \\ | ? * or end with a dot')
    if title.lower() in WINDOWS_RESERVED or title.startswith("_"):
        raise ApiError(400, f"{title!r} cannot be used as a plan title")
    return title


def read_plan(name):
    real = find_plan(name)
    if not real:
        raise ApiError(404, f"No plan named {name!r}")
    days = {d: {"session": "", "items": []} for d in DAYS}
    header, rows = read_plan_rows(plan_path(real))
    col = {h.lower(): i for i, h in enumerate(header)}
    get = lambda row, h: row[col[h.lower()]] if h.lower() in col else None
    for row in rows:
        day = _str(row[0])[:3].title()
        if day not in days:
            continue
        if _str(row[1]) and not days[day]["session"]:
            days[day]["session"] = _str(row[1])
        if _str(row[2]):
            lo, hi = parse_amount(get(row, "Amount"))
            days[day]["items"].append(
                {
                    "exercise": _str(row[2]),
                    "sets": _num(row[3]),
                    "amountMin": lo,
                    "amountMax": hi,
                    "unit": _unit(get(row, "Unit")),
                    "perSide": _yes(get(row, "Per Side")),
                    "load": _num(get(row, "Load")),
                    "loadType": _load_type(get(row, "Load Type")),
                    "emom": _yes(get(row, "EMOM")),
                    "hidden": _yes(get(row, "Hidden")),
                    "notes": _str(get(row, "Notes")),
                }
            )
    return {"name": real, "days": days}


def days_to_rows(days):
    for d in DAYS:
        day = days.get(d) or {}
        session = _str(day.get("session"))
        items = day.get("items") or []
        if not items and session:
            yield (d, session) + (None,) * (len(PLAN_HEADERS) - 2)
        for it in items:
            load_type = _load_type(it.get("loadType"))
            yield (
                d,
                session,
                _str(it.get("exercise")),
                _num(it.get("sets")),
                format_amount(_num(it.get("amountMin")), _num(it.get("amountMax"))),
                _unit(it.get("unit")),
                "yes" if it.get("perSide") else None,
                _num(it.get("load")) if load_type in ("kg", "+kg", "-kg") else None,
                load_type,
                "yes" if it.get("emom") else None,
                "yes" if it.get("hidden") else None,
                _str(it.get("notes")),
            )


def save_plan(title, days, replace=None):
    title = validate_title(title)
    existing = find_plan(title)
    if replace:
        old = find_plan(replace)
        if not old:
            raise ApiError(404, f"No plan named {replace!r}")
        if existing and existing.lower() != old.lower():
            raise ApiError(409, f"A plan named {title!r} already exists")
        if old != title:
            locked(lambda: plan_path(old).rename(plan_path(title)))
    if days is not None:  # None = rename only, keep the file's contents
        locked(lambda: write_plan_file(plan_path(title), list(days_to_rows(days))))
    return title


def delete_plan(name):
    """Moves the plan file to Plans/_deleted (never erases it)."""
    real = find_plan(name)
    if not real:
        raise ApiError(404, f"No plan named {name!r}")
    if len(plan_names()) == 1:
        raise ApiError(400, "Cannot delete the last remaining plan")
    DELETED_DIR.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y-%m-%d %H%M%S")
    locked(lambda: shutil.move(str(plan_path(real)), DELETED_DIR / f"{real} ({stamp}).xlsx"))


# ---------------------------------------------------------------- exports
def safe_export_name(name):
    name = _str(name)
    if not name or any(c in BAD_NAME_CHARS or ord(c) < 32 for c in name) or not name.lower().endswith((".png", ".pdf")):
        raise ApiError(400, f"Bad export file name {name!r}")
    return PLANS_DIR / name


def find_browser():
    """A Chromium-based browser (Edge / Chrome / Chromium) for headless PDF printing, or None."""
    if env := os.environ.get("FITVIEW_BROWSER"):
        return env if Path(env).exists() or shutil.which(env) else None
    candidates = {
        "Windows": [
            r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
            r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
            r"C:\Program Files\Google\Chrome\Application\chrome.exe",
            r"C:\Program Files (x86)\Google\Chrome\Application\chrome.exe",
            str(Path(os.environ.get("LOCALAPPDATA", "")) / "Google/Chrome/Application/chrome.exe"),
        ],
        "Darwin": [
            "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome",
            "/Applications/Microsoft Edge.app/Contents/MacOS/Microsoft Edge",
            "/Applications/Chromium.app/Contents/MacOS/Chromium",
            "/Applications/Brave Browser.app/Contents/MacOS/Brave Browser",
        ],
    }.get(platform.system(), [])
    for path in candidates:
        if path and Path(path).exists():
            return path
    for name in ("google-chrome", "google-chrome-stable", "chromium", "chromium-browser", "microsoft-edge", "msedge", "chrome"):
        if found := shutil.which(name):
            return found
    return None


def export_pdf(plan_name, days=None):
    """Prints /print to PDF with a headless Chromium browser. days = the on-screen (possibly unsaved) plan."""
    real = find_plan(plan_name)
    if not real:
        raise ApiError(404, f"No plan named {plan_name!r}")
    browser = find_browser()
    if not browser:
        raise ApiError(500, "PDF export needs Google Chrome, Microsoft Edge or Chromium (or set FITVIEW_BROWSER)")
    out = safe_export_name(f"{real} - Overview.pdf")
    if out.exists():
        locked(lambda: open(out, "ab").close())  # fail early if a PDF viewer has it open
    token = uuid.uuid4().hex
    if days is not None:
        PRINT_DAYS[token] = days
    started = time.time()
    try:
        # Reused profile dir: the browser's crash handler can keep it busy after exit, so a temp dir can't be cleaned up.
        profile = Path(tempfile.gettempdir()) / "fitview-browser-profile"
        subprocess.run(
            [
                browser,
                "--headless=new",
                "--disable-gpu",
                "--no-first-run",
                "--no-default-browser-check",
                f"--user-data-dir={profile}",
                "--no-pdf-header-footer",
                "--virtual-time-budget=5000",
                f"--print-to-pdf={out}",
                f"http://{HOST}:{PORT}/print?plan={quote(real)}&token={token}",
            ],
            capture_output=True,
            timeout=90,
        )
        # Some browsers (msedge.exe) hand off to a background process and return early: wait for a finished file.
        last_size = -1
        for _ in range(120):
            if out.exists() and out.stat().st_mtime >= started - 1:
                size = out.stat().st_size
                if size > 0 and size == last_size:
                    return out
                last_size = size
            time.sleep(0.5)
        raise ApiError(500, "The browser did not produce a PDF")
    finally:
        PRINT_DAYS.pop(token, None)


def print_page(plan_name, token=""):
    muscles, exercises, _ = read_library()
    plan = read_plan(plan_name)
    if token in PRINT_DAYS:
        plan["days"] = PRINT_DAYS[token]
    data = {"plan": plan, "muscles": muscles, "exercises": exercises, "generated": datetime.now().strftime("%d %b %Y"), "settings": read_settings()}
    html = (STATIC / "print.html").read_text(encoding="utf-8")
    payload = json.dumps(data).replace("</", "<\\/")
    return html.replace("/*DATA*/null", payload).encode("utf-8")


def open_folder(path):
    system = platform.system()
    if system == "Windows":
        os.startfile(path)  # noqa: S606 - local desktop app
    elif system == "Darwin":
        subprocess.Popen(["open", str(path)])
    else:
        subprocess.Popen(["xdg-open", str(path)])


# ---------------------------------------------------------------- http
class Handler(BaseHTTPRequestHandler):
    def log_message(self, fmt, *args):
        pass

    def _send(self, status, body, ctype="application/json"):
        data = body if isinstance(body, bytes) else json.dumps(body).encode()
        self.send_response(status)
        self.send_header("Content-Type", ctype)
        self.send_header("Content-Length", str(len(data)))
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        self.wfile.write(data)

    def _body(self):
        n = int(self.headers.get("Content-Length") or 0)
        return json.loads(self.rfile.read(n) or b"{}")

    def _handle(self, fn, lock=True):
        try:
            if lock:
                with LOCK:
                    result = fn()
            else:
                result = fn()
            self._send(200, result)
        except ApiError as e:
            self._send(e.status, {"error": str(e)})
        except Exception as e:  # surface unexpected errors in the UI instead of a dead request
            self._send(500, {"error": f"{type(e).__name__}: {e}"})

    def do_GET(self):
        url = urlparse(self.path)
        query = parse_qs(url.query)
        if url.path == "/api/data":

            def fn():
                muscles, exercises, unknown = read_library()
                return {
                    "muscles": muscles,
                    "exercises": exercises,
                    "plans": plan_names(),
                    "warnings": unknown,
                    "settings": read_settings(),
                    "version": __version__,
                    "workspace": str(HOME),
                }

            return self._handle(fn)
        if url.path == "/api/plan":
            return self._handle(lambda: read_plan(query.get("name", [""])[0]))
        if url.path == "/print":
            try:
                body = print_page(query.get("plan", [""])[0], query.get("token", [""])[0])
                return self._send(200, body, "text/html; charset=utf-8")
            except ApiError as e:
                return self._send(e.status, str(e).encode(), "text/plain")

        rel = "index.html" if url.path in ("/", "") else url.path.lstrip("/")
        f = (STATIC / rel).resolve()
        if STATIC.resolve() not in f.parents or not f.is_file():
            return self._send(404, b"Not found", "text/plain")
        self._send(200, f.read_bytes(), mimetypes.guess_type(f.name)[0] or "application/octet-stream")

    def do_POST(self):
        path = urlparse(self.path).path
        body = self._body()
        if path == "/api/plan/save":

            def fn():
                name = save_plan(body.get("name"), body.get("days"), replace=_str(body.get("replace")) or None)
                return {"ok": True, "plans": plan_names(), "name": name}

            return self._handle(fn)
        if path == "/api/plan/create":

            def fn():
                title = validate_title(body.get("name"))
                if find_plan(title):
                    raise ApiError(409, f"A plan named {title!r} already exists")
                src = _str(body.get("copyFrom"))
                days = body.get("days") or (read_plan(src)["days"] if src else {})
                save_plan(title, days)
                return {"ok": True, "plans": plan_names(), "name": title}

            return self._handle(fn)
        if path == "/api/settings":
            return self._handle(lambda: {"ok": True, "settings": write_settings(body)})
        if path == "/api/plan/delete":

            def fn():
                delete_plan(body.get("name"))
                return {"ok": True, "plans": plan_names()}

            return self._handle(fn)
        if path == "/api/export/pdf":
            # No lock: the headless browser calls back into /print while this runs.
            return self._handle(lambda: {"ok": True, "file": export_pdf(body.get("name"), body.get("days")).name}, lock=False)
        if path == "/api/export/images":

            def fn():
                saved = []
                for f in body.get("files") or []:
                    out = safe_export_name(f.get("name"))
                    data = base64.b64decode(_str(f.get("data")).split(",", 1)[-1])
                    locked(lambda out=out, data=data: out.write_bytes(data))
                    saved.append(out.name)
                return {"ok": True, "files": saved}

            return self._handle(fn)
        if path == "/api/open-folder":

            def fn():
                open_folder(PLANS_DIR)
                return {"ok": True}

            return self._handle(fn)
        self._send(404, {"error": "Unknown endpoint"})


def make_server(port, tries=10):
    """Binds to the first free port starting at `port`."""
    for p in range(port, port + tries):
        try:
            return ThreadingHTTPServer((HOST, p), Handler), p
        except OSError:
            continue
    raise SystemExit(f"No free port between {port} and {port + tries - 1} - pass --port to choose another one")


def main(argv=None):
    global PORT
    parser = argparse.ArgumentParser(prog="fitview", description="FitView - weekly training planner with a muscle body map.")
    parser.add_argument("--workspace", default=str(DEFAULT_HOME), help="folder for your library workbook and plans (default: ./workspace or $FITVIEW_HOME)")
    parser.add_argument("--port", type=int, default=DEFAULT_PORT, help=f"port to listen on (default {DEFAULT_PORT})")
    parser.add_argument("--no-browser", action="store_true", help="don't open the browser automatically")
    parser.add_argument("--version", action="version", version=f"FitView {__version__}")
    args = parser.parse_args(argv)

    configure(Path(args.workspace).expanduser().resolve())
    try:
        created = ensure_workspace(HOME)
    except LibraryError as e:
        raise SystemExit(str(e)) from None
    for item in created:
        print(f"Created {HOME / item}")

    server, PORT = make_server(args.port)
    url = f"http://{HOST}:{PORT}/"
    print(f"FitView {__version__} running at {url}  (Ctrl+C to stop)")
    print(f"Workspace: {HOME}")
    if not args.no_browser:
        threading.Timer(0.6, lambda: webbrowser.open(url)).start()
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        pass
    return 0


if __name__ == "__main__":
    sys.exit(main())
