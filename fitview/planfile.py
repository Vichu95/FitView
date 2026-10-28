"""Plan files: one weekly plan = one .xlsx file in the Plans folder, with a single sheet named "Plan".

Amounts are stored as numbers only ("10" or "8-12"); the unit, per-side flag and load type live in
their own columns so they are picked in the app instead of typed as free text.
"""

import re

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

DAYS = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]
PLAN_HEADERS = ["Day", "Session", "Exercise", "Sets", "Amount", "Unit", "Per Side", "Load", "Load Type", "EMOM", "Hidden", "Notes"]
UNITS = ["reps", "s", "min", "m", "steps", "rounds"]
LOAD_TYPES = ["kg", "+kg", "-kg", "band"]  # blank = bodyweight / none
NCOLS = len(PLAN_HEADERS)

_NUM = r"\d+(?:\.\d+)?"


def _n(s):
    f = float(s)
    return int(f) if f.is_integer() else f


def parse_amount(value):
    """'10' / '8-12' / 10 -> (min, max or None). Anything else -> (None, None)."""
    if value is None or value == "":
        return None, None
    if isinstance(value, (int, float)):
        return _n(value), None
    m = re.fullmatch(rf"\s*({_NUM})\s*(?:[-–]\s*({_NUM}))?\s*", str(value))
    if not m:
        return None, None
    lo, hi = _n(m.group(1)), (_n(m.group(2)) if m.group(2) else None)
    return lo, (hi if hi is not None and hi != lo else None)


def format_amount(lo, hi):
    if lo is None:
        return None
    return lo if hi is None else f"{lo}-{hi}"


def write_plan_file(path, rows):
    """rows: iterable of rows in PLAN_HEADERS order."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Plan"
    ws.append(PLAN_HEADERS)
    for r in rows:
        ws.append([v if v not in ("", None, False) else None for v in r])
    for c in range(1, NCOLS + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = PatternFill("solid", fgColor="1F2937")
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    for letter, width in zip("ABCDEFGHIJKL", (7, 24, 44, 7, 9, 8, 9, 8, 10, 7, 8, 40)):
        ws.column_dimensions[letter].width = width
    lists = (('"' + ",".join(UNITS) + '"', "F"), ('"yes"', "G"), ('"' + ",".join(LOAD_TYPES) + '"', "I"), ('"yes"', "J"), ('"yes"', "K"))
    for formula, col in lists:
        dv = DataValidation(type="list", formula1=formula, allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}500")
    wb.save(path)


def read_plan_rows(path):
    """Returns (header, rows) from the first sheet of a plan file; rows padded to the header width."""
    wb = load_workbook(path, read_only=True)
    try:
        it = wb.worksheets[0].iter_rows(values_only=True)
        header = [str(h or "").strip() for h in next(it, [])]
        width = max(len(header), NCOLS)
        rows = [list(r) + [None] * (width - len(r)) for r in it]
        return header, rows
    finally:
        wb.close()
